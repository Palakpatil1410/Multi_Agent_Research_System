import streamlit as st
import time
import datetime
import re
import base64
import io

# ── Page config ───────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="ResearchMind · AI Research System",
    page_icon="🔬",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ── Session state bootstrap ───────────────────────────────────────────────────
for key, default in {
    "history":     [],    # list of {id, topic, ts, state}
    "next_id":     1,
    "active_id":   None,
    "_last_state": None,
    "_last_topic": None,
}.items():
    if key not in st.session_state:
        st.session_state[key] = default

# ── CSS ───────────────────────────────────────────────────────────────────────
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Space+Grotesk:wght@300;400;500;600;700&family=DM+Mono:ital,wght@0,300;0,400;1,300&display=swap');

html, body, [class*="css"] { font-family: 'Space Grotesk', sans-serif; }
.stApp { background: #0C0E14; color: #E8EAF0; }
#MainMenu, footer, header { visibility: hidden; }
.block-container { padding: 2rem 3rem 4rem; max-width: 960px; }

[data-testid="stSidebar"] { background: #0E1018 !important; border-right: 1px solid #1E2233 !important; }
[data-testid="stSidebar"] * { color: #B0B5CC !important; }

.stTextInput input, .stTextArea textarea {
    background: #141720 !important; border: 1px solid #252A3A !important;
    border-radius: 10px !important; color: #E8EAF0 !important;
    font-family: 'Space Grotesk', sans-serif !important;
    font-size: 0.95rem !important; padding: 0.75rem 1rem !important;
    caret-color: #5B6AFF;
}
.stTextInput input:focus, .stTextArea textarea:focus {
    border-color: #5B6AFF !important;
    box-shadow: 0 0 0 3px rgba(91,106,255,0.15) !important;
}

.stButton > button {
    background: #5B6AFF !important; color: #fff !important;
    border: none !important; border-radius: 8px !important;
    font-family: 'Space Grotesk', sans-serif !important;
    font-size: 0.95rem !important; font-weight: 600 !important;
    padding: 0.65rem 1.5rem !important;
    transition: background 0.2s, transform 0.1s !important; width: 100%;
}
.stButton > button:hover  { background: #4858E8 !important; transform: translateY(-1px) !important; }
.stButton > button:active { transform: translateY(0) !important; }

.pipeline-grid { display: grid; grid-template-columns: repeat(4, 1fr); gap: 1rem; margin: 2rem 0; }
.step-card { background: #141720; border: 1px solid #252A3A; border-radius: 10px; padding: 1.1rem 1rem; transition: border-color 0.3s; }
.step-card.active { border-color: #5B6AFF; box-shadow: 0 0 20px rgba(91,106,255,0.12); }
.step-card.done   { border-color: #22C55E; }
.step-card.error  { border-color: #EF4444; }
.step-icon  { font-size: 1.4rem; margin-bottom: 0.5rem; }
.step-label { font-family: 'DM Mono', monospace; font-size: 0.65rem; letter-spacing: 0.12em; color: #4A5070; text-transform: uppercase; margin-bottom: 0.2rem; }
.step-name  { font-size: 0.9rem; font-weight: 600; color: #C8CBDF; }
.step-status { font-size: 0.75rem; color: #4A5070; margin-top: 0.3rem; }
.step-status.running { color: #5B6AFF; }
.step-status.done    { color: #22C55E; }
.step-status.error   { color: #EF4444; }

.result-block { background: #141720; border: 1px solid #252A3A; border-radius: 12px; padding: 1.5rem 1.75rem; margin-bottom: 1.25rem; }
.result-block-header { display: flex; align-items: center; gap: 0.6rem; margin-bottom: 1rem; }
.result-block-icon  { font-size: 1.1rem; }
.result-block-title { font-size: 0.8rem; font-family: 'DM Mono', monospace; letter-spacing: 0.12em; text-transform: uppercase; color: #5B6AFF; }
.result-block-body  { font-size: 0.92rem; line-height: 1.75; color: #B0B5CC; white-space: pre-wrap; word-break: break-word; }

.rm-wordmark { font-family: 'DM Mono', monospace; font-size: 0.75rem; letter-spacing: 0.18em; color: #5B6AFF; text-transform: uppercase; margin-bottom: 0.25rem; }
.rm-headline { font-size: 3rem; font-weight: 700; line-height: 1.1; letter-spacing: -0.03em; color: #F2F4FF; margin: 0 0 0.5rem; }
.rm-subline  { font-size: 1rem; font-weight: 300; color: #6E7490; margin-bottom: 2.5rem; }
.rm-divider  { border: none; border-top: 1px solid #1E2233; margin: 2rem 0; }
.err-banner  { background: #1A0A0A; border: 1px solid #7F1D1D; border-radius: 8px; padding: 1rem 1.25rem; color: #FCA5A5; font-size: 0.9rem; margin-top: 1rem; }
.topic-tag   { display: inline-block; background: rgba(91,106,255,0.12); color: #8A96FF; font-family: 'DM Mono', monospace; font-size: 0.78rem; padding: 0.3rem 0.8rem; border-radius: 999px; margin-bottom: 1.5rem; letter-spacing: 0.06em; }
</style>
""", unsafe_allow_html=True)


# ═══════════════════════════════════════════════════════════════
#  HELPERS
# ═══════════════════════════════════════════════════════════════
def to_str(value) -> str:
    if isinstance(value, str): return value
    if isinstance(value, list):
        return "\n".join(item.get("text", str(item)) if isinstance(item, dict) else str(item) for item in value)
    return str(value)

def truncate(value, limit: int = 2000) -> str:
    s = to_str(value)
    return s[:limit] + ("…" if len(s) > limit else "")

def result_block(icon, title, body):
    return (f'<div class="result-block"><div class="result-block-header">'
            f'<span class="result-block-icon">{icon}</span>'
            f'<span class="result-block-title">{title}</span></div>'
            f'<div class="result-block-body">{body}</div></div>')

def render_pipeline(active=-1, done_up_to=-1, errored=-1):
    STEPS = [("🔍","01","Search Agent"),("📄","02","Reader Agent"),("✍️","03","Writer Chain"),("🧐","04","Critic Chain")]
    html = '<div class="pipeline-grid">'
    for i,(icon,num,name) in enumerate(STEPS):
        if i==errored:       cls,badge,bcls = "error","failed","error"
        elif i==active:      cls,badge,bcls = "active","running…","running"
        elif i<=done_up_to:  cls,badge,bcls = "done","done ✓","done"
        else:                cls,badge,bcls = "","waiting",""
        html += (f'<div class="step-card {cls}"><div class="step-icon">{icon}</div>'
                 f'<div class="step-label">Step {num}</div><div class="step-name">{name}</div>'
                 f'<div class="step-status {bcls}">{badge}</div></div>')
    return html + "</div>"

def make_link(label, data, mime, filename, disabled=False, hint=""):
    if disabled:
        return (f'<a style="display:inline-flex;align-items:center;justify-content:center;width:100%;'
                f'padding:0.6rem 0;border-radius:8px;border:1px solid #252A3A;background:#0C0E14;'
                f'color:#3D4260;font-size:0.9rem;font-weight:600;text-decoration:none;cursor:not-allowed;"'
                f' title="{hint}">{label}</a>')
    b64  = base64.b64encode(data.encode("utf-8") if isinstance(data, str) else data).decode()
    href = f"data:{mime};base64,{b64}"
    return (f'<a href="{href}" download="{filename}" style="display:inline-flex;align-items:center;'
            f'justify-content:center;width:100%;padding:0.6rem 0;border-radius:8px;border:none;'
            f'background:#5B6AFF;color:#fff;font-size:0.9rem;font-weight:600;text-decoration:none;"'
            f' onmouseover="this.style.background=\'#4858E8\'" onmouseout="this.style.background=\'#5B6AFF\'">'
            f'{label}</a>')

def build_docx(topic, ss, sc, rp, fb):
    try:
        from docx import Document as D
        from docx.shared import Pt, RGBColor, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        doc = D()
        for sec in doc.sections:
            sec.top_margin = sec.bottom_margin = Inches(1)
            sec.left_margin = sec.right_margin = Inches(1.2)
        def ah(txt, lvl=1):
            p = doc.add_heading(txt, level=lvl)
            if p.runs: p.runs[0].font.color.rgb = RGBColor(0x1E,0x29,0x9E)
        def ab(txt):
            for ln in txt.split('\n'): doc.add_paragraph(ln if ln.strip() else "")
        tp = doc.add_heading("Research Report", level=0)
        tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if tp.runs: tp.runs[0].font.color.rgb = RGBColor(0x0C,0x0E,0x40)
        sp = doc.add_paragraph(topic); sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if sp.runs:
            sp.runs[0].font.size = Pt(13)
            sp.runs[0].font.color.rgb = RGBColor(0x5B,0x6A,0xFF)
        dp = doc.add_paragraph(datetime.date.today().strftime("Generated on %B %d, %Y"))
        dp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph("")
        for h,c in [("🔍 Search Results",ss),("📄 Scraped Content",sc),("✍️ Report",rp),("🧐 Critic Feedback",fb)]:
            ah(h); ab(c); doc.add_paragraph("")
        buf = io.BytesIO(); doc.save(buf); buf.seek(0)
        return buf.read(), None
    except ImportError: return None, "pip install python-docx"
    except Exception as ex: return None, str(ex)

def render_downloads(topic, state):
    st_ = topic[:40].replace(' ','_').replace('/','-')
    ss,sc,rp,fb = to_str(state['search_results']),to_str(state['scrape_content']),to_str(state['report']),to_str(state['feedback'])
    div = "\n"+"="*60+"\n"
    md = (f"# Research Report: {topic}\n\n---\n\n## 🔍 Search Results\n\n{ss}\n\n---\n\n"
          f"## 📄 Scraped Content\n\n{sc}\n\n---\n\n## ✍️ Report\n\n{rp}\n\n---\n\n## 🧐 Critic Feedback\n\n{fb}\n")
    tx = (f"RESEARCH REPORT: {topic.upper()}\n{'='*60}\n\nSEARCH RESULTS\n{div}{ss}\n\n"
          f"SCRAPED CONTENT\n{div}{sc}\n\nREPORT\n{div}{rp}\n\nCRITIC FEEDBACK\n{div}{fb}\n")
    db,de = build_docx(topic,ss,sc,rp,fb)
    l1 = make_link("📝 Markdown (.md)",   md, "text/markdown", f"research_{st_}.md")
    l2 = make_link("📃 Plain Text (.txt)", tx, "text/plain",    f"research_{st_}.txt")
    l3 = (make_link("📄 Word Doc (.docx)", db,
          "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
          f"research_{st_}.docx") if db else
          make_link("📄 Word Doc (.docx)", b"","","", disabled=True, hint=de))
    st.markdown(f"""
    <div style="margin-top:1.5rem;margin-bottom:0.75rem;">
        <span style="font-family:'DM Mono',monospace;font-size:0.72rem;letter-spacing:0.14em;
                     text-transform:uppercase;color:#5B6AFF;">⬇ Download Report</span>
    </div>
    <div style="display:grid;grid-template-columns:1fr 1fr 1fr;gap:0.75rem;">{l1}{l2}{l3}</div>
    """, unsafe_allow_html=True)

def render_results(state):
    st.markdown(
        result_block("🔍","Search Results",   truncate(state["search_results"]))
        + result_block("📄","Scraped Content", truncate(state["scrape_content"]))
        + result_block("✍️","Draft Report",    truncate(state["report"], 3000))
        + result_block("🧐","Critic Feedback", to_str(state["feedback"])),
        unsafe_allow_html=True,
    )


# ═══════════════════════════════════════════════════════════════
#  SIDEBAR — History
# ═══════════════════════════════════════════════════════════════
with st.sidebar:
    st.markdown(
        '<div style="font-family:DM Mono,monospace;font-size:0.7rem;letter-spacing:0.16em;'
        'color:#5B6AFF;text-transform:uppercase;padding-bottom:0.5rem;'
        'border-bottom:1px solid #1E2233;margin-bottom:1rem;">🕘 Research History</div>',
        unsafe_allow_html=True,
    )

    history = st.session_state.history

    if not history:
        st.markdown(
            '<p style="font-size:0.82rem;color:#3D4260;font-style:italic;">'
            'No sessions yet.<br>Run a topic to save it here.</p>',
            unsafe_allow_html=True,
        )
    else:
        if st.button("🗑  Clear all history", use_container_width=True):
            st.session_state.history     = []
            st.session_state.active_id   = None
            st.session_state._last_state = None
            st.session_state._last_topic = None
            st.rerun()

        st.markdown("<br>", unsafe_allow_html=True)

        for entry in reversed(history):
            is_active = st.session_state.active_id == entry["id"]
            label = ("▶ " if is_active else "") + entry["topic"][:30] + ("…" if len(entry["topic"]) > 30 else "")
            col_btn, col_del = st.sidebar.columns([5, 1])
            with col_btn:
                if st.button(label, key=f"h_{entry['id']}", use_container_width=True):
                    st.session_state.active_id   = entry["id"]
                    st.session_state._last_state = None
                    st.session_state._last_topic = None
                    st.rerun()
            with col_del:
                if st.button("✕", key=f"del_{entry['id']}"):
                    st.session_state.history = [e for e in st.session_state.history if e["id"] != entry["id"]]
                    if st.session_state.active_id == entry["id"]:
                        st.session_state.active_id = None
                    st.rerun()
            st.markdown(
                f'<div style="font-family:DM Mono,monospace;font-size:0.62rem;color:#3D4260;'
                f'margin:-0.3rem 0 0.6rem 0.2rem;">{entry["ts"]}</div>',
                unsafe_allow_html=True,
            )


# ═══════════════════════════════════════════════════════════════
#  MAIN
# ═══════════════════════════════════════════════════════════════
st.markdown('<div class="rm-wordmark">⬡ ResearchMind</div>', unsafe_allow_html=True)
st.markdown('<h1 class="rm-headline">Deep Research,<br>Automated.</h1>', unsafe_allow_html=True)
st.markdown('<p class="rm-subline">Multi-agent AI pipeline — search, scrape, write, critique — on any topic.</p>', unsafe_allow_html=True)
st.markdown('<hr class="rm-divider">', unsafe_allow_html=True)

# ── Viewing a saved history entry ─────────────────────────────────────────────
if st.session_state.active_id is not None:
    entry = next((e for e in st.session_state.history if e["id"] == st.session_state.active_id), None)
    if entry:
        col_back, col_ts = st.columns([1, 5])
        with col_back:
            if st.button("← New Research"):
                st.session_state.active_id = None
                st.rerun()
        with col_ts:
            st.markdown(
                f'<span style="color:#4A5070;font-family:DM Mono,monospace;font-size:0.78rem;">'
                f'{entry["ts"]}</span>', unsafe_allow_html=True)
        st.markdown(f'<div class="topic-tag">Topic · {entry["topic"]}</div>', unsafe_allow_html=True)
        render_results(entry["state"])
        render_downloads(entry["topic"], entry["state"])
        st.stop()

# ── Show results after pipeline completes ────────────────────────────────────
if st.session_state._last_state and st.session_state.active_id is None:
    last_topic = st.session_state._last_topic or ""
    last_state = st.session_state._last_state
    st.success("✓ Research complete — saved to history (see sidebar)!")
    st.markdown(f'<div class="topic-tag">Topic · {last_topic}</div>', unsafe_allow_html=True)
    render_results(last_state)
    render_downloads(last_topic, last_state)
    if st.button("▶  Start New Research", key="btn_new"):
        st.session_state._last_state = None
        st.session_state._last_topic = None
        st.rerun()
    st.stop()

# ── New research form ─────────────────────────────────────────────────────────
topic = st.text_area(
    "Research topic",
    placeholder="e.g.  Recent advances in quantum error correction (2024–2025)",
    height=90,
    label_visibility="collapsed",
)
col_btn, col_hint = st.columns([1, 5])
with col_btn:
    run_btn = st.button("▶  Run Research", use_container_width=True)
with col_hint:
    st.markdown(
        '<span style="color:#3D4260;font-size:0.82rem;line-height:3.2;">'
        'Runs 4 sequential agents · saved to history automatically</span>',
        unsafe_allow_html=True,
    )

# ── Pipeline execution ────────────────────────────────────────────────────────
if run_btn:
    if not topic.strip():
        st.warning("Please enter a research topic before running.")
        st.stop()

    st.markdown('<hr class="rm-divider">', unsafe_allow_html=True)
    st.markdown(f'<div class="topic-tag">Topic · {topic.strip()}</div>', unsafe_allow_html=True)

    pipeline_ph = st.empty()
    status_ph   = st.empty()   # live step message
    timer_ph    = st.empty()   # elapsed time
    results_ph  = st.empty()
    pipeline_ph.markdown(render_pipeline(), unsafe_allow_html=True)

    start_time = time.time()
    state = {}

    def update_timer():
        elapsed = int(time.time() - start_time)
        timer_ph.markdown(
            f'<div style="font-family:DM Mono,monospace;font-size:0.72rem;'
            f'color:#3D4260;margin-bottom:0.5rem;">⏱ {elapsed}s elapsed</div>',
            unsafe_allow_html=True,
        )

    try:
        from agents import invoke_search_agent, invoke_reader_agent, invoke_writer, invoke_critic

        # Step 1 — Search
        pipeline_ph.markdown(render_pipeline(active=0), unsafe_allow_html=True)
        status_ph.markdown(
            '<div style="color:#5B6AFF;font-size:0.88rem;margin-bottom:0.5rem;">'
            '🔍 Searching the web for sources…</div>', unsafe_allow_html=True)
        update_timer()
        sr = invoke_search_agent(topic)
        state["search_results"] = sr["messages"][-1].content
        update_timer()
        pipeline_ph.markdown(render_pipeline(active=1, done_up_to=0), unsafe_allow_html=True)
        status_ph.markdown(
            '<div style="color:#5B6AFF;font-size:0.88rem;margin-bottom:0.5rem;">'
            '📄 Scraping top URLs for deeper content…</div>', unsafe_allow_html=True)
        results_ph.markdown(result_block("🔍","Search Results",truncate(state["search_results"])), unsafe_allow_html=True)

        # Step 2 — Reader
        rr = invoke_reader_agent(topic, to_str(state["search_results"]))
        state["scrape_content"] = rr["messages"][-1].content
        update_timer()
        pipeline_ph.markdown(render_pipeline(active=2, done_up_to=1), unsafe_allow_html=True)
        status_ph.markdown(
            '<div style="color:#5B6AFF;font-size:0.88rem;margin-bottom:0.5rem;">'
            '✍️ Writing the research report…</div>', unsafe_allow_html=True)
        results_ph.markdown(
            result_block("🔍","Search Results",truncate(state["search_results"]))
            + result_block("📄","Scraped Content",truncate(state["scrape_content"])),
            unsafe_allow_html=True)

        # Step 3 — Writer
        state["report"] = invoke_writer(topic,
            f"SEARCH RESULTS:\n{to_str(state['search_results'])}\n\nDETAILED SCRAPE CONTENT:\n{to_str(state['scrape_content'])}")
        update_timer()
        pipeline_ph.markdown(render_pipeline(active=3, done_up_to=2), unsafe_allow_html=True)
        status_ph.markdown(
            '<div style="color:#5B6AFF;font-size:0.88rem;margin-bottom:0.5rem;">'
            '🧐 Critic is reviewing the report…</div>', unsafe_allow_html=True)
        results_ph.markdown(
            result_block("🔍","Search Results",truncate(state["search_results"]))
            + result_block("📄","Scraped Content",truncate(state["scrape_content"]))
            + result_block("✍️","Draft Report",truncate(state["report"],3000)),
            unsafe_allow_html=True)

        # Step 4 — Critic
        state["feedback"] = invoke_critic(state["report"])
        elapsed = int(time.time() - start_time)
        pipeline_ph.markdown(render_pipeline(done_up_to=3), unsafe_allow_html=True)
        status_ph.markdown(
            f'<div style="color:#22C55E;font-size:0.88rem;margin-bottom:0.5rem;">'
            f'✓ All done in {elapsed}s</div>', unsafe_allow_html=True)
        timer_ph.empty()
        results_ph.markdown(
            result_block("🔍","Search Results",truncate(state["search_results"]))
            + result_block("📄","Scraped Content",truncate(state["scrape_content"]))
            + result_block("✍️","Draft Report",truncate(state["report"],3000))
            + result_block("🧐","Critic Feedback",to_str(state["feedback"])),
            unsafe_allow_html=True)

        # Save to session history then rerun so sidebar updates
        st.session_state.history.append({
            "id":    st.session_state.next_id,
            "topic": topic.strip(),
            "ts":    datetime.datetime.now().strftime("%b %d, %Y  %H:%M"),
            "state": state,
        })
        st.session_state.next_id    += 1
        st.session_state._last_state = state
        st.session_state._last_topic = topic.strip()
        st.rerun()

    except Exception as e:
        err_str       = str(e)
        is_rate_limit = "429" in err_str or "RESOURCE_EXHAUSTED" in err_str or "quota" in err_str.lower()
        pipeline_ph.markdown(render_pipeline(errored=0), unsafe_allow_html=True)

        if is_rate_limit:
            dm       = re.search(r"retry[^\d]*(\d+)", err_str, re.IGNORECASE)
            rsecs    = int(dm.group(1)) if dm else 20
            is_daily = "PerDay" in err_str or "free_tier" in err_str.lower()
            if is_daily:
                st.markdown(f"""<div class="err-banner"><strong>🚫 Daily quota exhausted</strong><br><br>
                    You've used all <strong>20 free requests/day</strong>.<br><br>
                    <strong>Options:</strong><br>
                    &nbsp;• Wait until tomorrow<br>
                    &nbsp;• Add billing at <a href="https://ai.dev/rate-limit" target="_blank" style="color:#8A96FF;">ai.dev/rate-limit</a><br>
                    &nbsp;• Switch to <code>gemini-2.0-flash</code> in <code>agents.py</code>
                </div>""", unsafe_allow_html=True)
            else:
                st.markdown(f'<div class="err-banner"><strong>⏳ Rate limit — wait {rsecs}s then retry.</strong></div>', unsafe_allow_html=True)
                ph = st.empty()
                for r in range(rsecs, 0, -1):
                    ph.markdown(f"<div style='text-align:center;font-family:DM Mono,monospace;font-size:2rem;color:#5B6AFF;margin:1rem 0;'>Retry in {r}s…</div>", unsafe_allow_html=True)
                    time.sleep(1)
                ph.markdown("<div style='text-align:center;font-size:1.2rem;color:#22C55E;margin:1rem 0;'>✓ Ready — click Run Research</div>", unsafe_allow_html=True)
        else:
            st.markdown(f'<div class="err-banner"><strong>Pipeline error:</strong> {e}</div>', unsafe_allow_html=True)
            st.exception(e)